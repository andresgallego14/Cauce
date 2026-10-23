import openpyxl 
from infraestructura.database import Base, engine, SessionLocal
from infraestructura.models import Registro
from dominio.enums import Fuente

Base.metadata.create_all(engine)
db = SessionLocal()

ARCHIVO = "Xlsx/Consumos de Agua.xlsx"
HOJAS = {
    "Fertilizada 2026 " : Fuente.RIEGO,
    "Cruda 2026" : Fuente.CRUDA,
}
USUARIO_HISTORICO = "Brayan Castañeda"

wb = openpyxl.load_workbook(ARCHIVO, data_only=True)
total_importados = 0 

for nombre_hoja, fuente in HOJAS.items():
    ws = wb[nombre_hoja]
    filas = list(ws.iter_rows(min_row=8, max_row=ws.max_row))
    lectura_anterior = None
    for fila in filas:
        fecha = fila[10].value
        lectura = fila[11].value
        if fecha is None:
            break
        if not isinstance(lectura, (int, float)):
            continue
        if lectura_anterior is None:
            lectura_anterior = lectura
            continue
        valor = lectura - lectura_anterior
        lectura_anterior = lectura 
        if valor < 0:
            continue
        nuevo_registro = Registro(fecha=fecha.date(), tipo=fuente.tipo, fuente=fuente, valor=float(valor), lectura_acumulada=float(lectura), usuario=USUARIO_HISTORICO, notas="importado desde Excel historico")
        db.add(nuevo_registro)
        total_importados += 1
db.commit()
db.close()
print(f"listo. se importaron {total_importados} registros")

