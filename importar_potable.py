import openpyxl
from infraestructura.database import SessionLocal,Base,engine
from infraestructura.models import Registro
from dominio.enums import Fuente

Base.metadata.create_all(engine)
db = SessionLocal()
ARCHIVOS = {
    "Xlsx/REGISTRO DE CONSUMO  DE AGUA POTABLE (Respuestas).xlsx" : Fuente.AGUA_POTABLE_370,
    "Xlsx/_REGISTRO DE CONSUMO  DE AGUA POTABLE 60 (Respuestas).xlsx" : Fuente.AGUA_POTABLE_60
}

HOJA = "Respuestas de formulario 1"
USUARIO_HISTORICO = "Estefania Rendon"
total_importadas = 0

for nombre_archivo, fuente in ARCHIVOS.items():
    wb = openpyxl.load_workbook(nombre_archivo, data_only=True) 
    ws = wb[HOJA]
    


    ultima_lectura = None 
    for fila in ws.iter_rows(min_row=2, max_row=ws.max_row):
        marca_temporal = fila[0].value
        blancos_texto = fila[3].value
        if marca_temporal is None:
            continue
        try:
            lectura = float(blancos_texto)
        except (ValueError,TypeError):
            continue
        ultima_lectura = lectura

    if ultima_lectura is not None:
        nuevo = Registro(fecha = marca_temporal.date(), tipo = fuente.tipo, fuente = fuente, valor = 0.0, lectura_acumulada = ultima_lectura, usuario = USUARIO_HISTORICO, notas = "Importado desde forms")
        db.add(nuevo)
        total_importadas +=1
db.commit()
db.close()
print(f"Se importaron {total_importadas} de agua potable")









