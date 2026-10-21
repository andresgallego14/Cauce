import openpyxl
from database import Base, engine, SessionLocal
from models import Registro
from dominio.enums import Fuente


Base.metadata.create_all(engine)
db = SessionLocal()

ARCHIVO = "Xlsx/REGISTRO DE CONSUMO DE AGUA FERTIRRIEGO (Respuestas).xlsx"
HOJA = "Respuestas de formulario 1"
USUARIO_HISTORICO = "Estefania Rendon"
FUENTE = Fuente.FERTIRRIEGO

wb = openpyxl.load_workbook(ARCHIVO, data_only= True)
ws = wb[HOJA]
total_impotados = 0
lectura_anterior = None

for fila in ws.iter_rows(min_row= 54, max_row=ws.max_row):

    marca_temporal = fila[0].value
    semana_texto = fila [2].value
    metros_cubicos = fila [3].value
    if marca_temporal is None or not isinstance (metros_cubicos, (float, int)):
        continue
    if not isinstance (semana_texto, str) or not semana_texto. strip():
        continue
    if lectura_anterior is None:
        lectura_anterior = metros_cubicos
        continue
    valor = metros_cubicos - lectura_anterior
    lectura_anterior = metros_cubicos
    if valor < 0:
        continue
    nuevo_resgistro = Registro(fecha = marca_temporal.date(), tipo = FUENTE.tipo, fuente = FUENTE,
                                bocatoma = None , lectura_acumulada = float(metros_cubicos), valor = float(valor),
                                usuario = USUARIO_HISTORICO, notas = "importado desde forms"  ) 
    db.add(nuevo_resgistro)
    total_impotados += 1
db.commit()
db.close()
print(f"Se importaron datos de fertirriego {total_impotados}")



