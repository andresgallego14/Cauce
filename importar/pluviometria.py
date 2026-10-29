from datetime import date, datetime, timedelta

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from dominio.enums import Fuente
from importar.comun import reemplazar_registros
from infraestructura.models import Registro

HOJA = "Respuestas de formulario 1"
FUENTE = Fuente.LLUVIA
NOTAS = "importa desde google forms - pluviometria"
FILA_INICIAL = 2
COLUMNA_MARCA = 1
COLUMNA_DIA = 5
COLUMNA_MM = 6
DIAS_SEMANA = {
    "LUNES": 0,
    "MARTES": 1,
    "MIERCOLES": 2,
    "JUEVES": 3,
    "VIERNES": 4,
    "SABADO": 5,
    "DOMINGO": 6,
}


def fecha_del_dato(marca_temporal: datetime, dia_texto: object) -> date | None:
    if not isinstance(dia_texto, str) or not dia_texto.strip():
        return None
    nombre_dia = dia_texto.split()[-1].upper()
    if nombre_dia not in DIAS_SEMANA:
        return None
    diferencia = (marca_temporal.weekday() - DIAS_SEMANA[nombre_dia]) % 7
    return marca_temporal.date() - timedelta(days=diferencia)


def leer_lluvias(hoja: Worksheet) -> list[tuple[date, float]]:
    ultimas = {}
    for fila in range(FILA_INICIAL, hoja.max_row + 1):
        marca = hoja.cell(row=fila, column=COLUMNA_MARCA).value
        dia_texto = hoja.cell(row=fila, column=COLUMNA_DIA).value
        mm = hoja.cell(row=fila, column=COLUMNA_MM).value
        if marca is None or not isinstance(mm, (int, float)):
            continue
        fecha = fecha_del_dato(marca, dia_texto)
        if fecha is None:
            continue
        if fecha not in ultimas or marca > ultimas[fecha][0]:
            ultimas[fecha] = (marca, mm)
    return [(fecha, mm) for fecha, (marca, mm) in sorted(ultimas.items())]


def armar_registros_lluvia(
    lluvias: list[tuple[date, float]], usuario: str
) -> list[Registro]:
    return [
        Registro(
            fecha=fecha,
            tipo=FUENTE.tipo,
            fuente=FUENTE,
            valor=float(mm),
            usuario=usuario,
            notas=NOTAS,
        )
        for fecha, mm in lluvias
    ]


def importar_pluviometria(db: Session, libro: Workbook, usuario: str) -> int:
    lluvias = leer_lluvias(libro[HOJA])
    registros = armar_registros_lluvia(lluvias, usuario)
    return reemplazar_registros(db, FUENTE, NOTAS, registros)
