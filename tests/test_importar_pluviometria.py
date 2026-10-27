from datetime import date, datetime
import pytest
from openpyxl import Workbook
from ayudas import guardar_movimiento
from dominio.enums import Fuente
from importar.pluviometria import (
    NOTAS,
    armar_registros_lluvia,
    fecha_del_dato,
    importar_pluviometria,
    leer_lluvias,
)
from infraestructura.models import Registro

HOJA = "Respuestas de formulario 1"
LUNES = datetime(2026, 10, 26, 9, 0)
LUNES_MAS_TARDE = datetime(2026, 10, 26, 11, 0)
MARTES = datetime(2026, 10, 27, 8, 30)
ENVIOS = [(LUNES, "Domingo", 2.0), (LUNES_MAS_TARDE, "Lunes", 0.0)]
COLUMNA_MARCA = 1
COLUMNA_DIA = 5
COLUMNA_MM = 6

def escribir_fila(hoja, fila, marca, dia, mm):
    hoja.cell(row=fila, column=COLUMNA_MARCA, value=marca)
    hoja.cell(row=fila, column=COLUMNA_DIA, value=dia)
    hoja.cell(row=fila, column=COLUMNA_MM, value=mm)

def llenar_hoja(hoja, filas):
    for numero, (marca, dia, mm) in enumerate(filas):
        escribir_fila(hoja, 2 + numero, marca, dia, mm)

def hoja_con(filas):
    hoja = Workbook().active
    llenar_hoja(hoja, filas)
    return hoja

def libro_con(filas):
    libro = Workbook()
    libro.active.title = HOJA
    llenar_hoja(libro[HOJA], filas)
    return libro

def test_el_mismo_dia_de_la_semana_es_la_fecha_del_envio():
    assert fecha_del_dato(MARTES, "Martes") == date(2026, 10, 27)

def test_un_dia_anterior_retrocede_hasta_ese_dia():
    assert fecha_del_dato(MARTES, "Lunes") == date(2026, 10, 26)

def test_un_dia_posterior_retrocede_a_la_semana_anterior():
    assert fecha_del_dato(MARTES, "Sabado") == date(2026, 10, 24)

def test_el_dato_del_domingo_enviado_el_lunes_queda_en_domingo():
    assert fecha_del_dato(LUNES, "Domingo") == date(2026, 10, 25)

def test_el_dato_del_lunes_enviado_el_lunes_queda_en_lunes():
    assert fecha_del_dato(LUNES, "Lunes") == date(2026, 10, 26)

def test_toma_la_ultima_palabra_sin_importar_mayusculas():
    assert fecha_del_dato(MARTES, "dia lunes") == date(2026, 10, 26)

@pytest.mark.parametrize("dia_texto", ["Festivo", "   ", None])
def test_devuelve_none_si_no_reconoce_el_dia(dia_texto):
    assert fecha_del_dato(MARTES, dia_texto) is None

def test_lee_la_fecha_real_y_los_mm_desde_la_fila_2():
    hoja = hoja_con([(MARTES, "Martes", 5.5)])
    assert leer_lluvias(hoja) == [(date(2026, 10, 27), 5.5)]

def test_ignora_el_encabezado_de_la_fila_1():
    hoja = hoja_con([(MARTES, "Martes", 5.5)])
    escribir_fila(hoja, 1, MARTES, "Martes", 99.0)
    assert leer_lluvias(hoja) == [(date(2026, 10, 27), 5.5)]

def test_un_dia_sin_lluvia_se_guarda_como_cero():
    hoja = hoja_con([(MARTES, "Martes", 0)])
    assert leer_lluvias(hoja) == [(date(2026, 10, 27), 0)]

def test_el_dato_del_domingo_y_el_del_lunes_se_guardan_los_dos():
    hoja = hoja_con(ENVIOS)
    assert leer_lluvias(hoja) == [
        (date(2026, 10, 25), 2.0),
        (date(2026, 10, 26), 0.0),
    ]

def test_si_hay_dos_envios_para_el_mismo_dia_gana_el_mas_reciente():
    hoja = hoja_con([(LUNES_MAS_TARDE, "Lunes", 3.0), (LUNES, "Lunes", 1.0)])
    assert leer_lluvias(hoja) == [(date(2026, 10, 26), 3.0)]

def test_salta_las_filas_malas():
    hoja = hoja_con(
        [
            (None, "Martes", 5.0),
            (MARTES, "Martes", "sin dato"),
            (MARTES, "Festivo", 5.0),
            (MARTES, None, 5.0),
            (MARTES, "Martes", 4.0),
        ]
    )
    assert leer_lluvias(hoja) == [(date(2026, 10, 27), 4.0)]

def test_arma_un_registro_por_cada_lluvia():
    lluvias = [(date(2026, 10, 26), 3.0), (date(2026, 10, 27), 0)]
    registros = armar_registros_lluvia(lluvias, "prueba")
    assert [registro.fecha for registro in registros] == [
        date(2026, 10, 26),
        date(2026, 10, 27),
    ]
    assert [registro.valor for registro in registros] == [3.0, 0.0]


def test_el_registro_de_lluvia_completa_fuente_tipo_usuario_y_notas():
    registro = armar_registros_lluvia([(date(2026, 10, 26), 3.0)], "prueba")[0]
    assert registro.fuente == Fuente.LLUVIA
    assert registro.tipo == Fuente.LLUVIA.tipo
    assert registro.usuario == "prueba"
    assert registro.notas == NOTAS
    assert registro.lectura_acumulada is None

def test_importa_cada_lluvia_como_fuente_lluvia(db):
    total = importar_pluviometria(db, libro_con(ENVIOS), "prueba")
    registros = db.query(Registro).all()
    assert total == 2
    assert sorted(registro.valor for registro in registros) == [0.0, 2.0]
    assert all(registro.fuente == Fuente.LLUVIA for registro in registros)

def test_importar_dos_veces_no_duplica(db):
    for _ in range(2):
        importar_pluviometria(db, libro_con(ENVIOS), "prueba")
    assert db.query(Registro).count() == 2

def test_no_borra_lo_cargado_desde_la_app(db):
    guardar_movimiento(
        db, date(2026, 10, 20), Fuente.LLUVIA, 7.0, "registrado en la app"
    )
    db.commit()
    importar_pluviometria(db, libro_con(ENVIOS), "prueba")
    assert db.query(Registro).count() == 3