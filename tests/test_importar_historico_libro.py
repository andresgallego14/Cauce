from datetime import date, datetime

from openpyxl import Workbook

from ayudas import guardar_movimiento
from dominio.enums import Fuente
from importar.historico import importar_historico
from infraestructura.models import Registro

HOJA_RIEGO = "Fertilizada 2026 "
HOJA_CRUDA = "Cruda 2026"
RIEGO = [(datetime(2026, 1, 1), 100), (datetime(2026, 1, 2), 110)]
CRUDA = [(datetime(2026, 1, 1), 50), (datetime(2026, 1, 2), 70)]


def escribir(hoja, lecturas):
    for numero, (fecha, lectura) in enumerate(lecturas):
        hoja.cell(row=8 + numero, column=11, value=fecha)
        hoja.cell(row=8 + numero, column=12, value=lectura)


def libro_con(riego, cruda):
    libro = Workbook()
    hoja_riego = libro.active
    hoja_riego.title = HOJA_RIEGO
    escribir(hoja_riego, riego)
    escribir(libro.create_sheet(HOJA_CRUDA), cruda)
    return libro


def test_importa_cada_hoja_con_su_fuente(db):
    total = importar_historico(db, libro_con(RIEGO, CRUDA), "prueba")
    riego = db.query(Registro).filter(Registro.fuente == Fuente.RIEGO).all()
    cruda = db.query(Registro).filter(Registro.fuente == Fuente.CRUDA).all()
    assert total == 2
    assert [registro.valor for registro in riego] == [10.0]
    assert [registro.valor for registro in cruda] == [20.0]


def test_importar_dos_veces_no_duplica(db):
    for _ in range(2):
        importar_historico(db, libro_con(RIEGO, CRUDA), "prueba")
    assert db.query(Registro).count() == 2


def test_no_borra_lo_cargado_desde_la_app(db):
    guardar_movimiento(db, date(2026, 1, 1), Fuente.RIEGO, 7.0, "registrado en la app")
    db.commit()
    importar_historico(db, libro_con(RIEGO, CRUDA), "prueba")
    assert db.query(Registro).count() == 3
