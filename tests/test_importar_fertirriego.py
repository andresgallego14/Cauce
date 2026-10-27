from datetime import date, datetime

from openpyxl import Workbook

from ayudas import guardar_movimiento
from dominio.enums import Fuente
from importar.fertirriego import importar_fertirriego, leer_lecturas
from infraestructura.models import Registro

HOJA = "Respuestas de formulario 1"
COLUMNA_FECHA = 1
COLUMNA_SEMANA = 3
COLUMNA_LECTURA = 4
FILAS = [
    (datetime(2026, 1, 1), "semana 1", 100),
    (datetime(2026, 1, 2), "semana 2", 110),
    (datetime(2026, 1, 3), "semana 3", 125),
]


def escribir_fila(hoja, fila, fecha, semana, lectura):
    hoja.cell(row=fila, column=COLUMNA_FECHA, value=fecha)
    hoja.cell(row=fila, column=COLUMNA_SEMANA, value=semana)
    hoja.cell(row=fila, column=COLUMNA_LECTURA, value=lectura)


def llenar_hoja(hoja, filas):
    for numero, (fecha, semana, lectura) in enumerate(filas):
        escribir_fila(hoja, 54 + numero, fecha, semana, lectura)


def hoja_con(filas):
    hoja = Workbook().active
    llenar_hoja(hoja, filas)
    return hoja


def libro_con(filas):
    libro = Workbook()
    libro.active.title = HOJA
    llenar_hoja(libro[HOJA], filas)
    return libro


def test_lee_fecha_y_lectura_desde_la_fila_54():
    hoja = hoja_con(
        [
            (datetime(2026, 1, 1), "semana 1", 100),
            (datetime(2026, 1, 2), "semana 2", 110),
        ]
    )

    assert leer_lecturas(hoja) == [(date(2026, 1, 1), 100), (date(2026, 1, 2), 110)]


def test_ignora_lo_que_hay_encima_de_la_fila_54():
    hoja = hoja_con([(datetime(2026, 1, 2), "semana 2", 110)])
    escribir_fila(hoja, 53, datetime(2026, 1, 1), "semana 1", 100)

    assert leer_lecturas(hoja) == [(date(2026, 1, 2), 110)]


def test_salta_las_filas_sin_texto_de_semana():
    hoja = hoja_con(
        [
            (datetime(2026, 1, 1), None, 100),
            (datetime(2026, 1, 2), "   ", 105),
            (datetime(2026, 1, 3), "semana 3", 110),
        ]
    )

    assert leer_lecturas(hoja) == [(date(2026, 1, 3), 110)]


def test_salta_las_filas_malas_y_sigue_leyendo():
    hoja = hoja_con(
        [
            (datetime(2026, 1, 1), "semana 1", 100),
            (None, "semana 2", 105),
            (datetime(2026, 1, 3), "semana 3", "sin dato"),
            (datetime(2026, 1, 4), "semana 4", 120),
        ]
    )

    assert leer_lecturas(hoja) == [(date(2026, 1, 1), 100), (date(2026, 1, 4), 120)]


def test_importa_los_consumos_como_fertirriego(db):
    total = importar_fertirriego(db, libro_con(FILAS), "prueba")

    registros = db.query(Registro).all()
    assert total == 2
    assert sorted(registro.valor for registro in registros) == [10.0, 15.0]
    assert all(registro.fuente == Fuente.FERTIRRIEGO for registro in registros)
    assert all(registro.usuario == "prueba" for registro in registros)


def test_importar_dos_veces_no_duplica(db):
    for _ in range(2):
        importar_fertirriego(db, libro_con(FILAS), "prueba")

    assert db.query(Registro).count() == 2


def test_no_borra_lo_cargado_desde_la_app(db):
    guardar_movimiento(
        db, date(2026, 1, 1), Fuente.FERTIRRIEGO, 7.0, "registrado en la app"
    )
    db.commit()

    importar_fertirriego(db, libro_con(FILAS), "prueba")

    assert db.query(Registro).count() == 3