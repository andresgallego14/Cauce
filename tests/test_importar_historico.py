from datetime import date, datetime
from openpyxl import Workbook
from importar.historico import leer_lecturas

COLUMNA_FECHA = 11
COLUMNA_LECTURA = 12

def hoja_con(filas, fila_inicial=8):
    hoja = Workbook().active
    for numero, (fecha, lectura) in enumerate(filas):
        hoja.cell(row=fila_inicial + numero, column=COLUMNA_FECHA, value=fecha)
        hoja.cell(row=fila_inicial + numero, column=COLUMNA_LECTURA, value=lectura)
    return hoja

def test_lee_fecha_y_lectura_desde_la_fila_8():
    hoja = hoja_con([(datetime(2026, 1, 1), 100), (datetime(2026, 1, 2), 110)])
    assert leer_lecturas(hoja) == [(date(2026, 1, 1), 100), (date(2026, 1, 2), 110)]

def test_ignora_lo_que_hay_encima_de_la_fila_8():
    hoja = hoja_con([(datetime(2026, 1, 1), 100)])
    hoja.cell(row=7, column=COLUMNA_FECHA, value="Fecha")
    hoja.cell(row=7, column=COLUMNA_LECTURA, value="Lectura")
    assert leer_lecturas(hoja) == [(date(2026, 1, 1), 100)]

def test_se_detiene_en_la_primera_fecha_vacia():
    hoja = hoja_con(
        [
            (datetime(2026, 1, 1), 100),
            (None, None),
            (datetime(2026, 1, 3), 120),
        ]
    )
    assert leer_lecturas(hoja) == [(date(2026, 1, 1), 100)]

def test_salta_las_lecturas_que_no_son_numero():
    hoja = hoja_con(
        [
            (datetime(2026, 1, 1), 100),
            (datetime(2026, 1, 2), "sin dato"),
            (datetime(2026, 1, 3), 120),
        ]
    )
    assert leer_lecturas(hoja) == [(date(2026, 1, 1), 100), (date(2026, 1, 3), 120)]