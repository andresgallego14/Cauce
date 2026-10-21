import openpyxl
from database import Base, SessionLocal, engine
from models import Registro
from dominio.enums import Fuente
from datetime import timedelta

Base.metadata.create_all(engine)
db = SessionLocal()

ARCHIVO ="Xlsx/DATOS PLUVIOMETRÍA (Respuestas).xlsx"
HOJA = "Respuestas de formulario 1"
USUARIO_HISTORICO = "Estefania Rendon"
FUENTE = Fuente.LLUVIA
DIAS_SEMANA = {
    "LUNES" : 0,
    "MARTES" : 1,
    "MIERCOLES" : 2,
    "JUEVES" : 3,
    "VIERNES" : 4,
    "SABADO" : 5,
    "DOMINGO" : 6,
}
wb = openpyxl.load_workbook(ARCHIVO, data_only=True)
ws = wb[HOJA]

ultima_lectura_por_dia = {}

for fila in ws.iter_rows(min_row=2, max_row=ws.max_row):
    marca_temporal = fila[0].value
    dia_texto = fila[4].value
    mm = fila[5].value
    if marca_temporal is None or not isinstance(mm, (int, float)):
        continue
    if not isinstance (dia_texto, str) or not dia_texto.strip():
        continue

    nombre_dia = dia_texto.split()[-1].upper()
    if nombre_dia not in DIAS_SEMANA:
        continue 
    objetivo = DIAS_SEMANA[nombre_dia]
    diferencia = (marca_temporal.weekday() - objetivo) %7
    fecha_real = marca_temporal.date() - timedelta(days=diferencia)

    if fecha_real not in ultima_lectura_por_dia or marca_temporal > ultima_lectura_por_dia[fecha_real][0]:
        ultima_lectura_por_dia[fecha_real] = (marca_temporal, mm)

total_importados = 0
for fecha_real, (marca_temporal, mm) in ultima_lectura_por_dia.items():
        nuevo = Registro(
            fecha = fecha_real,
            tipo = FUENTE.tipo,
            fuente = FUENTE,
            valor = float(mm),
            bloques_id=None,
            usuario = USUARIO_HISTORICO,
            notas = "importa desde google forms - pluviometria",
        )
        db.add(nuevo)
        total_importados += 1
db.commit()
db.close()
print(f"se importaron los datos de pluviometria {total_importados}")